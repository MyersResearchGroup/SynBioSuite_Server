from __future__ import annotations 
from flask import Flask, render_template, request, jsonify
from flask_cors import CORS
from .main import app
import sys
import os
import json

import sbol2build
import tricahue
import sbol2
import pudu

import openpyxl
import pandas as pd

#routes
#check if the app is running
@app.route('/api/status')
def pin():
    return jsonify({"status": "working"}), 200

@app.route("/api/data")
def get_data():
    return app.send_static_file("data.json")

@app.route('/api/upload_sbs', methods = ['POST'])
def upload_file_from_sbs_post():
    if 'Metadata' not in request.files:
        print(request)
        return 'No file part', 400
    file = request.files['Metadata']
    if file.filename == '':
        return 'No selected file', 400
    file_contents = file.read()
    if 'Params' not in request.files:
        return 'No Params file part', 400
    params_file = request.files['Params']
    if params_file.filename == '':
        return 'No selected Params file', 400
    params_from_request = json.loads(params_file.read())

    # instantiate the XDC class using the params_from_request dictionary
    print(request.files['Metadata'])
    xdc = tricahue.XDC(input_excel_path = request.files['Metadata'],
            fj_url = params_from_request['fj_url'],
            fj_user = None, 
            fj_pass = None, 
            sbh_url = params_from_request['sbh_url'], 
            sbh_user = None, 
            sbh_pass = None, 
            sbh_collection = params_from_request['sbh_collec'], 
            sbh_collection_description = params_from_request['sbh_collec_desc'],
            sbh_overwrite = params_from_request['sbh_overwrite'], 
            fj_overwrite = params_from_request['fj_overwrite'], 
            fj_token = params_from_request['fj_token'], 
            sbh_token = params_from_request['sbh_token'],
            homespace = "https://synbiohub.org/gonza10v"
            )
            


    xdc.initialize()
    xdc.log_in_sbh()
    xdc.convert_to_sbol()
    xdc.generate_sbol_hash_map()

    try:
        sbh_url = xdc.upload_to_sbh()
    except AttributeError as e:
        return jsonify({"error": str(e)}), 400

    sbs_upload_response_dict ={
        "sbh_url": sbh_url,
        "status": "success"
    }
    return jsonify(sbs_upload_response_dict)

@app.route('/api/upload_sbs_up', methods = ['POST'])
def upload_file_from_sbs_post_up():
    if 'Metadata' not in request.files:
        print(request)
        return 'No file part', 400
    file = request.files['Metadata']
    if file.filename == '':
        return 'No selected file', 400
    file_contents = file.read()
    if 'Params' not in request.files:
        return 'No Params file part', 400
    params_file = request.files['Params']
    if params_file.filename == '':
        return 'No selected Params file', 400
    params_from_request = json.loads(params_file.read())

    # instantiate the XDC class using the params_from_request dictionary
    print(request.files['Metadata'])
    xdc = tricahue.XDC(input_excel_path = request.files['Metadata'],
            fj_url = params_from_request['fj_url'],
            fj_user = params_from_request['fj_user'],
            fj_pass = params_from_request['fj_pass'], 
            sbh_url = params_from_request['sbh_url'], 
            sbh_user = params_from_request['sbh_user'], 
            sbh_pass = params_from_request['sbh_pass'],
            sbh_collection = params_from_request['sbh_collec'], 
            sbh_collection_description = params_from_request['sbh_collec_desc'],
            sbh_overwrite = params_from_request['sbh_overwrite'], 
            fj_overwrite = params_from_request['fj_overwrite'], 
            fj_token = None, 
            sbh_token = None,
            homespace = "https://synbiohub.org/gonza10v"
            )
            


    xdc.initialize()
    xdc.log_in_sbh()
    xdc.convert_to_sbol()
    xdc.generate_sbol_hash_map()

    try:
        sbh_url = xdc.upload_to_sbh()
    except AttributeError as e:
        return jsonify({"error": str(e)}), 400
    
    sbs_upload_response_dict ={
        "sbh_url": sbh_url,
        "status": "success"
    }
    return jsonify(sbs_upload_response_dict)
    


@app.route('/sbol_2_build_golden_gate', methods=['POST'])
def sbol_2_build_golden_gate():
    # Error checking in the request
    print("request", request.files)

    if 'plasmid_backbone' not in request.files:
        return jsonify({"error": "Missing plasmid backbone"}), 400
    if 'insert_parts' not in request.files:
        return jsonify({"error": "Missing insert parts"}), 400
    if 'wizard_selections' not in request.form:
        return jsonify({"error": "Missing wizard selections"}), 400

    wizard_selection = request.form.get('wizard_selections')
    plasmid_backbone = request.files.get('plasmid_backbone')
    insert_parts = request.files.getlist('insert_parts')

    # Parse the json

    wizard_selection_json = json.loads(wizard_selection)
    assembly_method = wizard_selection_json.get('formValues').get('assemblyMethod')

    # Check if the assembly method is valid
    if assembly_method != 'MoClo':
        return jsonify({"error": "Invalid assembly method"}), 400
    
    # Get the restriction item
    restriction_enzyme = wizard_selection_json.get('formValues').get('restrictionEnzyme')

    # code for sbol2build
    part_docs = []
    for item in insert_parts:
        doc = sbol2.Document()
        doc.read(item)
        part_docs.append(doc)
    
    bb_doc = sbol2.Document()
    bb_doc.read(plasmid_backbone)

    assembly_doc = sbol2.Document()
    assembly_obj = sbol2build.golden_gate_assembly_plan('testassem', part_docs, bb_doc, restriction_enzyme, assembly_doc)

    try:
        composites = assembly_obj.run()

        return_string = assembly_doc.writeString()

        # Return the file as a response
        return return_string

    except ValueError as e:
        # catch sbol2build errors and return to frontend
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"Unexpected server error: {str(e)}"}), 500

@app.route('/build_pudu', methods=['POST'])
def build_pudu():
    # Error checking in the request
    print("request", request.files)

    if 'assembly_plan' not in request.files:
        return jsonify({"error": "Missing assembly plan"}), 400
    if 'wizard_selections' not in request.form:
        return jsonify({"error": "Missing wizard selections"}), 400

    wizard_selection = request.form.get('wizard_selections')
    assembly_plan_file = request.files.get('assembly_plan')

    # # Parse the json
    wizard_selection_json = json.loads(wizard_selection)
    build_method = wizard_selection_json.get('formValues').get('buildMethod')

    # Check if the assembly method is valid
    if build_method != 'PUDU':
        return jsonify({"error": "Invalid build method"}), 400
    
    # Get the assembly, is in text [?] and read using SBOL
    # transform assembly plan in text to a SBOL document, similar to the example following

    # part_docs = []
    # for item in insert_parts:
    #     doc = sbol2.Document()
    #     doc.read(item)
    #     part_docs.append(doc)

    try:
        assembly_plan = assembly_plan_file.read().decode('utf-8')
        assembly_plan_doc = sbol2.Document()
        assembly_plan_doc.readString(assembly_plan)
        # pudupy.dictionaryCreatorPython(assembly_plan_file)
    except Exception as e:
        return jsonify({"error": f"Error parsing file: {str(e)}"}), 400

    # TODO: transform assembly plan document to dictionary

    # TODO: use dictionary in PUDU

    try:
        # composites = assembly_obj.run()

        # return_string = assembly_doc.writeString()

        # # Return the file as a response
        # return return_string
        return jsonify({"message": "PUDU build not implemented yet"}), 501
    
    except ValueError as e:
        # catch errors and return to frontend
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"Unexpected server error: {str(e)}"}), 500


@app.route('/api/inspect_request', methods=['POST'])
def inspect_request():
    files = {}
    for name in request.files:
        file = request.files[name]
        try:
            files[name] = json.loads(file.read())
        except Exception as e:
            return jsonify({"error": str(e)}), 400

    return jsonify({
        "message": "Request received successfully", 
        "files": files}), 200

@app.route('/api/get_description_and_library_name', methods=['POST'])
def upload_excel():
    if 'file' not in request.files:
        return 'No file part', 400
    file = request.files['file']
    if file.filename == '':
        return 'No selected file', 400
    dataframe1 = pd.read_excel(file, engine='openpyxl')
    library_name = dataframe1[dataframe1.iloc[:, 1] == 'Library Name'].iloc[0, 2]
    description = dataframe1[dataframe1.iloc[:, 1] == 'Description'].iloc[0, 2]
    return jsonify({"name": library_name, "description": description})